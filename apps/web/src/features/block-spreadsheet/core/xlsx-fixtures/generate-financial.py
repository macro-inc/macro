"""Independent OOXML fixtures; no production/user data. Requires openpyxl.
Run from the repo root with python3 <this file>.
"""
from pathlib import Path
from datetime import datetime
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

out = Path(__file__).parent
w = Workbook()
a = w.active
a.title = 'Assumptions'
a.append(['Rate', .1])
a.append(['Principal', 250000])
a.append(['First date', datetime(2026, 1, 1)])
a.append(['Next date', datetime(2027, 1, 1)])
a['B1'].number_format = '0.0%'
a['B3'].number_format = a['B4'].number_format = 'yyyy-mm-dd'
a.sheet_state = 'hidden'
w.defined_names.add(DefinedName('DiscountRate', attr_text="Assumptions!$B$1"))
b = w.create_sheet('DCF model')
b['A1'] = 'Illustrative DCF — independent fixture'
b.merge_cells('A1:F1')
b.row_dimensions[1].height = 32
b.freeze_panes = 'B4'
b.row_dimensions[6].hidden = True
b.column_dimensions['F'].hidden = True
b.column_dimensions['A'].width = 28
b.append(['Cash flow', 'Year 0', 'Year 1', 'Year 2', 'Year 3'])
b.append(['Unlevered cash flow', -100, 100, 120, 140])
b['B4'] = '=NPV(DiscountRate,C3:E3)+B3'
b['B5'] = '=IRR(B7:C7)'
b['B7'], b['C7'] = -100, 110
b['B8'] = '=XIRR(B7:C7,Assumptions!B3:B4)'
b['B9'] = '=PMT(0.06/12,360,Assumptions!B2)'
b['B10'] = '=SUM($C$3:E3)'
b['C10'] = '=B10*DiscountRate'
b['B11'] = '=SUMIFS(C3:E3,C3:E3,">100")'
b['B12'] = '=INDEX(C3:E3,1,MATCH(120,C3:E3,0))'
b['B13'] = '=IFERROR(1/0,0)'
b['B14'] = '=TaxRate*100'
b.defined_names.add(DefinedName('TaxRate', attr_text='0.25'))
b['A16'] = '0000123'
b['B16'] = '=1/3'
b['C16'] = '=B16*3'
b['D16'] = 'literal =SUM(1,2)'
b['E16'] = True
for addr, value, fmt in [
 ('B18', -1250000.125, '#,##0.00;[Red](#,##0.00);"—"'),
 ('B19', 0, '#,##0.00;[Red](#,##0.00);"—"'),
 ('B20', 1250000, '0.0,,"m"'),
 ('B21', 1.25, '0.0"x"'),
 ('B22', .125, '0.00%'),
 ('B23', 1.5, '[h]:mm'),
 ('B24', 60, 'mm/dd/yyyy'),
]:
 b[addr] = value
 b[addr].number_format = fmt
b['B18'].font = Font(name='Calibri', size=11, color='0000FF')
b['B18'].border = Border(bottom=Side(style='double', color='123456'), top=Side(style='thin', color='000000'))
b['B18'].alignment = Alignment(horizontal='right', indent=1)
b.auto_filter.ref = 'A2:E3'
w.save(out / 'financial-model.xlsx')

# A namespaced workbook with absolute relationships, like the reported upload.
w = Workbook()
s = w.active
s.title = 'Companies'
s.append(['Company', 'Value'])
s.append(['Sample A', 42])
s.append(['Sample B', 58])
t = Table(displayName='CompaniesTable', ref='A1:B3')
t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
s.add_table(t)
s['D1'] = '=SUM(B2:B3)'
w.save(out / 'namespaced-table.xlsx')
with ZipFile(out / 'namespaced-table.xlsx') as z:
 files = {name: z.read(name) for name in z.namelist()}
for name, content in list(files.items()):
 if name.endswith(('.xml', '.rels')):
  root = ET.fromstring(content)
  if name.endswith('.rels'):
   base = name.rsplit('/_rels/', 1)[0] if '/_rels/' in name else ''
   import posixpath
   for rel in root:
    target = rel.get('Target', '')
    if rel.get('TargetMode') != 'External' and not target.startswith('/'):
     rel.set('Target', '/' + posixpath.normpath(posixpath.join(base, target)))
  files[name] = ET.tostring(root, encoding='utf-8', xml_declaration=True)
with ZipFile(out / 'namespaced-table.xlsx', 'w', ZIP_DEFLATED) as z:
 for name, data in files.items(): z.writestr(name, data)
