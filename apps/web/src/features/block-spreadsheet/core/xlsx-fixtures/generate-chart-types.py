"""Generate chart-types.xlsx: radar, bubble, stock and surface charts written by
openpyxl, not Macro.

Run from this directory: python3 generate-chart-types.py (needs openpyxl).
"""

import datetime

import openpyxl
from openpyxl.chart import (
    BubbleChart,
    RadarChart,
    Reference,
    Series,
    StockChart,
    SurfaceChart,
)
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.updown_bars import UpDownBars

workbook = openpyxl.Workbook()


def sized(chart, title, width=12):
    chart.title = title
    chart.width = width
    chart.height = 7.5
    return chart


# Scores of three products on five measures, as a radar.
ratings = workbook.active
ratings.title = "Ratings"
ratings.append(("Measure", "Basic", "Plus", "Pro"))
for row in [
    ("Price", 9, 6, 4),
    ("Speed", 4, 7, 9),
    ("Support", 5, 7, 8),
    ("Features", 3, 6, 9),
    ("Ease", 8, 7, 5),
]:
    ratings.append(row)
measures = Reference(ratings, min_col=1, min_row=2, max_row=6)

radar = sized(RadarChart(), "Product ratings")
radar.type = "marker"
radar.add_data(
    Reference(ratings, min_col=2, max_col=4, min_row=1, max_row=6),
    titles_from_data=True,
)
radar.set_categories(measures)
ratings.add_chart(radar, "F2")

filled = sized(RadarChart(), "Coverage")
filled.type = "filled"
filled.add_data(
    Reference(ratings, min_col=2, max_col=3, min_row=1, max_row=6),
    titles_from_data=True,
)
filled.set_categories(measures)
ratings.add_chart(filled, "F18")

# Stores by visits, average spend and floor area, as bubbles.
stores = workbook.create_sheet("Stores")
stores.append(("Visits", "Spend", "Area"))
for row in [(120, 35, 400), (300, 22, 900), (80, 60, 250), (210, 41, 1600)]:
    stores.append(row)
bubble = sized(BubbleChart(), "Stores")
bubble.series.append(
    Series(
        values=Reference(stores, min_col=2, min_row=2, max_row=5),
        xvalues=Reference(stores, min_col=1, min_row=2, max_row=5),
        zvalues=Reference(stores, min_col=3, min_row=2, max_row=5),
        title="Area",
    )
)
stores.add_chart(bubble, "E2")

# A week of share prices: open, high, low and close.
prices = workbook.create_sheet("Prices")
prices.append(("Date", "Open", "High", "Low", "Close"))
for index, (open_, high, low, close) in enumerate(
    [
        (102, 108, 100, 107),
        (107, 111, 104, 105),
        (105, 106, 98, 99),
        (99, 104, 97, 103),
        (103, 110, 102, 109),
    ]
):
    prices.append(
        (datetime.date(2024, 3, 4) + datetime.timedelta(days=index), open_, high, low, close)
    )
    prices.cell(row=2 + index, column=1).number_format = "mmm d"
days = Reference(prices, min_col=1, min_row=2, max_row=6)

ohlc = sized(StockChart(), "Open, high, low, close")
ohlc.add_data(
    Reference(prices, min_col=2, max_col=5, min_row=1, max_row=6),
    titles_from_data=True,
)
ohlc.set_categories(days)
for series in ohlc.series:
    series.graphicalProperties.line.noFill = True
ohlc.hiLowLines = ChartLines()
ohlc.upDownBars = UpDownBars(gapWidth=150)
prices.add_chart(ohlc, "G2")

hlc = sized(StockChart(), "High, low, close")
hlc.add_data(
    Reference(prices, min_col=3, max_col=5, min_row=1, max_row=6),
    titles_from_data=True,
)
hlc.set_categories(days)
for series in hlc.series:
    series.graphicalProperties.line.noFill = True
hlc.hiLowLines = ChartLines()
prices.add_chart(hlc, "G18")

# Yield by rainfall and temperature, as a contour.
yields = workbook.create_sheet("Yield")
yields.append(("Rain", "10°C", "15°C", "20°C", "25°C"))
for rain, row in [
    ("Low", (20, 35, 45, 30)),
    ("Medium", (30, 55, 70, 50)),
    ("High", (25, 50, 85, 60)),
]:
    yields.append((rain, *row))
contour = sized(SurfaceChart(), "Yield")
contour.add_data(
    Reference(yields, min_col=2, max_col=5, min_row=1, max_row=4),
    titles_from_data=True,
)
contour.set_categories(Reference(yields, min_col=1, min_row=2, max_row=4))
yields.add_chart(contour, "G2")

workbook.save("chart-types.xlsx")
