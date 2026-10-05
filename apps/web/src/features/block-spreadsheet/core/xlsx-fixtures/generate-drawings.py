"""Generate drawings.xlsx: charts and an image written by openpyxl, not Macro.

Run from this directory: python3 generate-drawings.py (needs openpyxl and Pillow).
"""

import io

import openpyxl
from openpyxl.chart import (
    AreaChart,
    BarChart,
    DoughnutChart,
    LineChart,
    PieChart,
    Reference,
    ScatterChart,
    Series,
)
from openpyxl.drawing.image import Image
from PIL import Image as PILImage
from PIL import ImageDraw

workbook = openpyxl.Workbook()
sales = workbook.active
sales.title = "Sales"
sales.append(("Month", "Revenue", "Costs", "Profit"))
for index, month in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun"]):
    revenue = 1000 + index * 250
    costs = 700 + index * 120
    sales.append((month, revenue, costs, revenue - costs))
sales["F1"] = "Region"
sales["G1"] = "Share"
for index, (region, share) in enumerate(
    [("North", 40), ("South", 25), ("East", 20), ("West", 15)]
):
    sales.cell(row=2 + index, column=6, value=region)
    sales.cell(row=2 + index, column=7, value=share)

months = Reference(sales, min_col=1, min_row=2, max_row=7)


def sized(chart, title, width=12):
    chart.title = title
    chart.width = width
    chart.height = 7
    return chart


column = sized(BarChart(), "Revenue and costs")
column.type = "col"
column.grouping = "clustered"
column.add_data(
    Reference(sales, min_col=2, min_row=1, max_col=3, max_row=7), titles_from_data=True
)
column.set_categories(months)
sales.add_chart(column, "I2")

line = sized(LineChart(), "Profit")
line.add_data(Reference(sales, min_col=4, min_row=1, max_row=7), titles_from_data=True)
line.set_categories(months)
sales.add_chart(line, "I17")

pie = sized(PieChart(), "Share by region", width=10)
pie.add_data(Reference(sales, min_col=7, min_row=1, max_row=5), titles_from_data=True)
pie.set_categories(Reference(sales, min_col=6, min_row=2, max_row=5))
sales.add_chart(pie, "Q2")

stacked = sized(BarChart(), "Stacked costs", width=10)
stacked.type = "bar"
stacked.grouping = "stacked"
stacked.overlap = 100
stacked.add_data(
    Reference(sales, min_col=2, min_row=1, max_col=3, max_row=7), titles_from_data=True
)
stacked.set_categories(months)
sales.add_chart(stacked, "Q17")

area = sized(AreaChart(), "Revenue area")
area.add_data(Reference(sales, min_col=2, min_row=1, max_row=7), titles_from_data=True)
area.set_categories(months)
sales.add_chart(area, "I32")

scatter = sized(ScatterChart(), "Revenue vs costs", width=10)
points = Series(
    Reference(sales, min_col=3, min_row=2, max_row=7),
    Reference(sales, min_col=2, min_row=2, max_row=7),
    title="Costs",
)
points.marker.symbol = "circle"
scatter.series.append(points)
sales.add_chart(scatter, "Q32")

# A plain logo without text, so the bytes are the same on every machine.
logo = PILImage.new("RGB", (160, 60), (30, 90, 160))
ImageDraw.Draw(logo).rectangle([10, 10, 150, 50], outline=(255, 255, 255), width=3)
data = io.BytesIO()
logo.save(data, format="PNG")
data.seek(0)
sales.add_image(Image(data), "A10")

# Another sheet's charts read Sales: a doughnut, and columns with a line on a
# secondary axis.
summary = workbook.create_sheet("Summary")
summary["A1"] = "Charts of the Sales sheet"
doughnut = sized(DoughnutChart(), "Regions", width=10)
doughnut.add_data(
    Reference(sales, min_col=7, min_row=1, max_row=5), titles_from_data=True
)
doughnut.set_categories(Reference(sales, min_col=6, min_row=2, max_row=5))
summary.add_chart(doughnut, "A3")

combination = sized(BarChart(), "Revenue and profit")
combination.add_data(
    Reference(sales, min_col=2, min_row=1, max_row=7), titles_from_data=True
)
combination.set_categories(months)
margin = LineChart()
margin.add_data(Reference(sales, min_col=4, min_row=1, max_row=7), titles_from_data=True)
margin.y_axis.axId = 200
margin.y_axis.crosses = "max"
combination += margin
summary.add_chart(combination, "H3")

workbook.save("drawings.xlsx")
